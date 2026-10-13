import asyncio
import json
import os
from pathlib import Path
import subprocess
import time
from types import SimpleNamespace
from unittest.mock import patch
import uuid
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.database import async_session
from app.llm.client import (
    ExtractedRequirement,
    FakeGeminiClient,
    GeminiClientError,
    GoogleGenAIClient,
    PageText,
    create_gemini_client,
)
from app.main import app
from app.models.evaluation import AnalysisJob, Document, DocumentPage, Evaluation
from app.services.ingestion import extract_docx_pages
from app.services.job_launcher import InlineJobLauncher, get_job_launcher
from app.services.worker import INTERRUPTED_MESSAGE, recover_interrupted_jobs, run_job_by_id
from tests.unit.test_documents_api import make_valid_pdf_bytes

client = TestClient(app)


def _fake_token(uid: str, email: str = "extract_user@vixenone.app") -> dict:
    return {"uid": uid, "email": email, "name": "Extract User"}


def _find_dataset_dir() -> Path:
    candidates = [
        Path("/app/dataset"),
        Path(__file__).resolve().parents[3] / "dataset",
        Path(__file__).resolve().parents[2] / "dataset",
    ]
    for c in candidates:
        if c.exists():
            return c
    return Path("/app/dataset")


def _make_real_docx_bytes() -> bytes:
    import io
    import docx

    d = docx.Document()
    d.add_heading("1. Overview", level=1)
    d.add_paragraph("Vendor Beta proposal text.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _setup_eval_with_all_docs(
    uid: str, pdf_pages: int = 2, v1_pdf_bytes: bytes | None = None
) -> tuple[str, dict, list[str]]:
    auth = {"Authorization": "Bearer token"}
    ev = client.post(
        "/api/evaluations",
        headers=auth,
        json={
            "title": "Andheri Housekeeping RFP",
            "category": "Housekeeping",
            "vendor_names": ["Vendor Alpha", "Vendor Beta"],
        },
    ).json()
    eval_id = ev["id"]
    vendor_ids = [v["id"] for v in ev["vendors"]]

    req_pdf = make_valid_pdf_bytes(page_count=pdf_pages)
    r_req = client.post(
        f"/api/evaluations/{eval_id}/documents",
        headers=auth,
        files={"file": ("rfp.pdf", req_pdf, "application/pdf")},
        data={"kind": "requirements"},
    )
    assert r_req.status_code == 200

    v1_pdf = v1_pdf_bytes if v1_pdf_bytes is not None else make_valid_pdf_bytes(page_count=2)
    r_v1 = client.post(
        f"/api/evaluations/{eval_id}/documents",
        headers=auth,
        files={"file": ("alpha.pdf", v1_pdf, "application/pdf")},
        data={"kind": "proposal", "vendor_id": vendor_ids[0]},
    )
    assert r_v1.status_code == 200

    v2_docx = _make_real_docx_bytes()
    r_v2 = client.post(
        f"/api/evaluations/{eval_id}/documents",
        headers=auth,
        files={
            "file": (
                "beta.docx",
                v2_docx,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"kind": "proposal", "vendor_id": vendor_ids[1]},
    )
    assert r_v2.status_code == 200

    return eval_id, r_req.json(), [r_v1.json()["id"], r_v2.json()["id"]]


def test_extraction_happy_path():
    uid = f"uid-ext-happy-{uuid.uuid4().hex}"
    fake_gemini = FakeGeminiClient(
        default_pages=[
            PageText(page_number=1, text="§1.1 Must deploy 12 operatives."),
            PageText(page_number=2, text="§2.1 Quarterly deep clean optional."),
        ],
        default_requirements=[
            ExtractedRequirement(
                code="",
                text="Deploy 12 housekeeping operatives daily",
                category="Staffing",
                mandatory=True,
                acceptance="12 staff on site",
                source_page=1,
                source_section="§1.1",
            ),
            ExtractedRequirement(
                code="REQ-02",
                text="Quarterly carpet deep cleaning",
                category="Service",
                mandatory=False,
                acceptance="Quarterly schedule",
                source_page=2,
                source_section="§2.1",
            ),
        ],
    )
    app.dependency_overrides[get_job_launcher] = lambda: InlineJobLauncher(
        async_session, lambda: fake_gemini
    )
    try:
        with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
            auth = {"Authorization": "Bearer token"}
            eval_id, req_doc, _ = _setup_eval_with_all_docs(uid, pdf_pages=2)

            resp = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
            assert resp.status_code == 202
            job_data = resp.json()
            assert job_data["kind"] == "extract_requirements"
            assert job_data["status"] == "done"
            assert job_data["finished_at"] is not None

            # Verify evaluation status is requirements_review and documents are ready
            detail = client.get(f"/api/evaluations/{eval_id}", headers=auth).json()
            assert detail["status"] == "requirements_review"
            assert detail["requirements_count"] == 2
            for doc in detail["documents"]:
                assert doc["status"] == "ready"
                assert doc["error_message"] is None

            # Verify jobs/latest returns timestamps and status
            latest_job = client.get(f"/api/evaluations/{eval_id}/jobs/latest", headers=auth).json()
            assert latest_job["status"] == "done"
            assert latest_job["started_at"] is not None
            assert latest_job["finished_at"] is not None
            assert latest_job["created_at"] is not None
            assert latest_job["updated_at"] is not None

            # Verify extracted requirements
            reqs = client.get(f"/api/evaluations/{eval_id}/requirements", headers=auth).json()
            assert len(reqs) == 2
            assert reqs[0]["code"] == "R1"  # missing code assigned R1
            assert reqs[0]["origin"] == "extracted"
            assert reqs[0]["source_document_id"] == req_doc["id"]
            assert reqs[0]["source_page"] == 1
            assert reqs[0]["position"] == 0
            assert reqs[1]["code"] == "REQ-02"
            assert reqs[1]["position"] == 1
    finally:
        app.dependency_overrides.pop(get_job_launcher, None)


def test_out_of_range_page_set_to_null():
    uid = f"uid-ext-page-range-{uuid.uuid4().hex}"
    fake_gemini = FakeGeminiClient(
        default_pages=[
            PageText(page_number=1, text="Page 1 text"),
            PageText(page_number=2, text="Page 2 text"),
        ],
        default_requirements=[
            ExtractedRequirement(
                code="",
                text="Valid page requirement",
                category="Service",
                mandatory=True,
                acceptance="Check",
                source_page=2,
                source_section="§1",
            ),
            ExtractedRequirement(
                code="",
                text="Out of range high page requirement",
                category="Service",
                mandatory=False,
                acceptance="Check",
                source_page=99,
                source_section="§2",
            ),
            ExtractedRequirement(
                code="",
                text="Out of range zero page requirement",
                category="Service",
                mandatory=False,
                acceptance="Check",
                source_page=0,
                source_section="§3",
            ),
        ],
    )
    app.dependency_overrides[get_job_launcher] = lambda: InlineJobLauncher(
        async_session, lambda: fake_gemini
    )
    try:
        with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
            auth = {"Authorization": "Bearer token"}
            eval_id, _, _ = _setup_eval_with_all_docs(uid, pdf_pages=2)
            resp = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
            assert resp.status_code == 202

            reqs = client.get(f"/api/evaluations/{eval_id}/requirements", headers=auth).json()
            assert len(reqs) == 3
            assert reqs[0]["code"] == "R1"
            assert reqs[0]["source_page"] == 2
            assert reqs[1]["code"] == "R2"
            assert reqs[1]["source_page"] is None
            assert reqs[2]["code"] == "R3"
            assert reqs[2]["source_page"] is None
    finally:
        app.dependency_overrides.pop(get_job_launcher, None)


def test_parallel_cap_respected():
    uid = f"uid-ext-cap-{uuid.uuid4().hex}"
    fake_gemini = FakeGeminiClient(delay_seconds=0.05)

    with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
        auth = {"Authorization": "Bearer token"}
        ev = client.post(
            "/api/evaluations",
            headers=auth,
            json={
                "title": "Cap Test RFP",
                "category": "Facilities",
                "vendor_names": ["V1", "V2", "V3"],
            },
        ).json()
        eval_id = ev["id"]
        vendor_ids = [v["id"] for v in ev["vendors"]]

        # Upload 1 requirements PDF + 3 vendor proposal PDFs
        pdf_bytes = make_valid_pdf_bytes(page_count=1)
        r_req = client.post(
            f"/api/evaluations/{eval_id}/documents",
            headers=auth,
            files={"file": ("req.pdf", pdf_bytes, "application/pdf")},
            data={"kind": "requirements"},
        ).json()
        for idx, vid in enumerate(vendor_ids):
            client.post(
                f"/api/evaluations/{eval_id}/documents",
                headers=auth,
                files={"file": (f"v{idx}.pdf", pdf_bytes, "application/pdf")},
                data={"kind": "proposal", "vendor_id": vid},
            )

        # Add 2 extra PDF documents in DB so total is 6 PDF documents (> cap 4)
        async def _seed_extra_docs_and_run():
            async with async_session() as s:
                for i in range(2):
                    s.add(
                        Document(
                            id=uuid.uuid4(),
                            evaluation_id=uuid.UUID(eval_id),
                            vendor_id=uuid.UUID(vendor_ids[0]),
                            kind="proposal",
                            filename=f"extra_{i}.pdf",
                            mime="application/pdf",
                            size_bytes=len(pdf_bytes),
                            storage_key=r_req["storage_key"],
                            page_count=1,
                            status="uploaded",
                        )
                    )
                job = AnalysisJob(
                    id=uuid.uuid4(),
                    evaluation_id=uuid.UUID(eval_id),
                    kind="extract_requirements",
                    status="queued",
                    step_label="Queued",
                )
                s.add(job)
                await s.commit()
                job_id = job.id

            await run_job_by_id(job_id, async_session, lambda: fake_gemini)

        asyncio.run(_seed_extra_docs_and_run())
        assert fake_gemini.pdf_call_count == 6
        assert 2 <= fake_gemini.max_in_flight <= 4


PROPOSAL_FAILURE_MARKER = b"FAIL_CORRUPTED_PROPOSAL"


def test_proposal_failure_job_done_and_doc_error():
    uid = f"uid-ext-prop-fail-{uuid.uuid4().hex}"

    def _pdf_handler(pdf_bytes: bytes) -> list[PageText]:
        # Fails strictly by file content (distinctive bytes marker), not call order
        if PROPOSAL_FAILURE_MARKER in pdf_bytes:
            raise GeminiClientError("Corrupted proposal PDF")
        return [PageText(page_number=1, text="Valid requirement page text")]

    fake_gemini = FakeGeminiClient(pdf_handler=_pdf_handler)
    app.dependency_overrides[get_job_launcher] = lambda: InlineJobLauncher(
        async_session, lambda: fake_gemini
    )
    try:
        with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
            auth = {"Authorization": "Bearer token"}
            corrupted_prop_pdf = make_valid_pdf_bytes(
                page_count=2, marker=PROPOSAL_FAILURE_MARKER
            )
            eval_id, req_doc, prop_doc_ids = _setup_eval_with_all_docs(
                uid, v1_pdf_bytes=corrupted_prop_pdf
            )

            resp = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
            assert resp.status_code == 202
            assert resp.json()["status"] == "done"

            detail = client.get(f"/api/evaluations/{eval_id}", headers=auth).json()
            assert detail["status"] == "requirements_review"
            docs_by_id = {d["id"]: d for d in detail["documents"]}
            assert docs_by_id[req_doc["id"]]["status"] == "ready"
            assert docs_by_id[prop_doc_ids[0]]["status"] == "error"
            assert "Corrupted proposal PDF" in docs_by_id[prop_doc_ids[0]]["error_message"]
            assert docs_by_id[prop_doc_ids[1]]["status"] == "ready"
    finally:
        app.dependency_overrides.pop(get_job_launcher, None)


def test_requirements_doc_failure_fails_job_and_preserves_status():
    uid = f"uid-ext-req-fail-{uuid.uuid4().hex}"

    def _pdf_handler(pdf_bytes: bytes) -> list[PageText]:
        raise GeminiClientError("Failed reading requirements PDF")

    fake_gemini = FakeGeminiClient(pdf_handler=_pdf_handler)
    app.dependency_overrides[get_job_launcher] = lambda: InlineJobLauncher(
        async_session, lambda: fake_gemini
    )
    try:
        with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
            auth = {"Authorization": "Bearer token"}
            eval_id, req_doc, _ = _setup_eval_with_all_docs(uid)

            resp = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
            assert resp.status_code == 202
            job_data = resp.json()
            assert job_data["status"] == "failed"
            assert "Requirements document extraction failed" in job_data["error_message"]

            detail = client.get(f"/api/evaluations/{eval_id}", headers=auth).json()
            assert detail["status"] == "draft"
            docs_by_id = {d["id"]: d for d in detail["documents"]}
            assert docs_by_id[req_doc["id"]]["status"] == "error"
    finally:
        app.dependency_overrides.pop(get_job_launcher, None)


def test_rate_limit_then_success_retried():
    calls = 0

    class RateLimitError(Exception):
        def __init__(self):
            super().__init__("429 Too Many Requests")
            self.status_code = 429

    class FakeModels:
        async def generate_content(self, model, contents, config):
            nonlocal calls
            calls += 1
            if calls < 3:
                raise RateLimitError()
            payload = {"pages": [{"page_number": 1, "text": "Extracted after retry"}]}
            return SimpleNamespace(text=json.dumps(payload))

    fake_sdk = SimpleNamespace(aio=SimpleNamespace(models=FakeModels()))
    gemini = GoogleGenAIClient(
        api_key="test-key",
        model="gemini-3.6-flash",
        base_backoff=0.01,
        sdk_client=fake_sdk,
    )

    pages = asyncio.run(gemini.pdf_to_pages(b"%PDF-1.4 dummy"))
    assert calls == 3
    assert len(pages) == 1
    assert pages[0].text == "Extracted after retry"


def test_invalid_json_twice_fails_with_message():
    calls = 0

    class FakeModels:
        async def generate_content(self, model, contents, config):
            nonlocal calls
            calls += 1
            return SimpleNamespace(text="{invalid json output")

    fake_sdk = SimpleNamespace(aio=SimpleNamespace(models=FakeModels()))
    gemini = GoogleGenAIClient(
        api_key="test-key",
        model="gemini-3.6-flash",
        base_backoff=0.01,
        sdk_client=fake_sdk,
    )

    with pytest.raises(GeminiClientError) as exc_info:
        asyncio.run(gemini.pdf_to_pages(b"%PDF-1.4 dummy"))
    assert calls == 2
    assert "invalid structured JSON" in str(exc_info.value)


def test_extract_422_missing_docs():
    uid = f"uid-ext-422-{uuid.uuid4().hex}"
    with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
        auth = {"Authorization": "Bearer token"}
        pdf_bytes = make_valid_pdf_bytes(1)

        # 1. Evaluation with only 1 vendor (and requirements + proposal uploaded) -> 422 "Add at least 2 vendors"
        ev_one = client.post(
            "/api/evaluations",
            headers=auth,
            json={"title": "One Vendor RFP", "category": "Facilities"},
        ).json()
        ev_one_id = ev_one["id"]
        v_only = client.post(
            f"/api/evaluations/{ev_one_id}/vendors",
            headers=auth,
            json={"name": "Solo Vendor"},
        ).json()
        client.post(
            f"/api/evaluations/{ev_one_id}/documents",
            headers=auth,
            files={"file": ("req.pdf", pdf_bytes, "application/pdf")},
            data={"kind": "requirements"},
        )
        client.post(
            f"/api/evaluations/{ev_one_id}/documents",
            headers=auth,
            files={"file": ("solo.pdf", pdf_bytes, "application/pdf")},
            data={"kind": "proposal", "vendor_id": v_only["id"]},
        )
        r_one_vendor = client.post(f"/api/evaluations/{ev_one_id}/extract", headers=auth)
        assert r_one_vendor.status_code == 422
        assert "Add at least 2 vendors" in r_one_vendor.json()["detail"]

        # 2. Evaluation with 2 vendors, missing requirements doc or vendor proposals
        ev = client.post(
            "/api/evaluations",
            headers=auth,
            json={
                "title": "Incomplete RFP",
                "category": "Facilities",
                "vendor_names": ["Vendor 1", "Vendor 2"],
            },
        ).json()
        eval_id = ev["id"]
        v1_id = ev["vendors"][0]["id"]

        # No docs at all -> 422
        r1 = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
        assert r1.status_code == 422
        assert "requirements document" in r1.json()["detail"].lower()

        # Upload requirements doc only (missing proposals) -> 422
        client.post(
            f"/api/evaluations/{eval_id}/documents",
            headers=auth,
            files={"file": ("req.pdf", pdf_bytes, "application/pdf")},
            data={"kind": "requirements"},
        )
        r2 = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
        assert r2.status_code == 422
        assert "Vendor 1" in r2.json()["detail"]
        assert "Vendor 2" in r2.json()["detail"]

        # Upload only Vendor 1 proposal (Vendor 2 proposal still missing) -> 422 naming Vendor 2
        client.post(
            f"/api/evaluations/{eval_id}/documents",
            headers=auth,
            files={"file": ("v1.pdf", pdf_bytes, "application/pdf")},
            data={"kind": "proposal", "vendor_id": v1_id},
        )
        r3 = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
        assert r3.status_code == 422
        assert r3.json()["detail"] == "Upload a proposal for Vendor 2 or remove the vendor."


def test_extract_409_active_job_and_upload_delete_409():
    uid = f"uid-ext-409-active-{uuid.uuid4().hex}"
    with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
        auth = {"Authorization": "Bearer token"}
        eval_id, req_doc, _ = _setup_eval_with_all_docs(uid)

        async def _insert_active_job():
            async with async_session() as s:
                job = AnalysisJob(
                    id=uuid.uuid4(),
                    evaluation_id=uuid.UUID(eval_id),
                    kind="extract_requirements",
                    status="extracting",
                    step_label="Reading documents 1 of 3",
                )
                s.add(job)
                await s.commit()

        asyncio.run(_insert_active_job())

        # Extract while active job -> 409
        r_ext = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
        assert r_ext.status_code == 409

        # Upload while active job -> 409
        pdf_bytes = make_valid_pdf_bytes(1)
        r_up = client.post(
            f"/api/evaluations/{eval_id}/documents",
            headers=auth,
            files={"file": ("new.pdf", pdf_bytes, "application/pdf")},
            data={"kind": "requirements"},
        )
        assert r_up.status_code == 409

        # Delete while active job -> 409
        r_del = client.delete(
            f"/api/evaluations/{eval_id}/documents/{req_doc['id']}",
            headers=auth,
        )
        assert r_del.status_code == 409


def test_extract_409_user_edited_without_confirm_and_success_with_confirm():
    uid = f"uid-ext-confirm-{uuid.uuid4().hex}"
    fake_gemini = FakeGeminiClient()
    app.dependency_overrides[get_job_launcher] = lambda: InlineJobLauncher(
        async_session, lambda: fake_gemini
    )
    try:
        with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
            auth = {"Authorization": "Bearer token"}
            eval_id, _, _ = _setup_eval_with_all_docs(uid)

            # Save a user-origin requirement
            put_res = client.put(
                f"/api/evaluations/{eval_id}/requirements",
                headers=auth,
                json=[
                    {
                        "code": "USER-1",
                        "text": "Custom user requirement",
                        "category": "Custom",
                        "mandatory": True,
                        "acceptance": "Custom acceptance",
                        "position": 0,
                        "origin": "user",
                    }
                ],
            )
            assert put_res.status_code == 200

            # Without confirm_replace -> 409 mentioning edits will be replaced
            r_no_confirm = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
            assert r_no_confirm.status_code == 409
            assert "replaced" in r_no_confirm.json()["detail"].lower()

            # With confirm_replace=True -> 202 and replaces with extracted requirements
            r_confirm = client.post(
                f"/api/evaluations/{eval_id}/extract",
                headers=auth,
                json={"confirm_replace": True},
            )
            assert r_confirm.status_code == 202
            reqs = client.get(f"/api/evaluations/{eval_id}/requirements", headers=auth).json()
            assert len(reqs) == 3
            assert all(r["origin"] == "extracted" for r in reqs)
    finally:
        app.dependency_overrides.pop(get_job_launcher, None)


def test_cross_user_404_on_extract():
    uid1 = f"uid-ext-u1-{uuid.uuid4().hex}"
    uid2 = f"uid-ext-u2-{uuid.uuid4().hex}"
    with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid1)):
        eval_id, _, _ = _setup_eval_with_all_docs(uid1)

    with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid2)):
        r = client.post(
            f"/api/evaluations/{eval_id}/extract",
            headers={"Authorization": "Bearer token2"},
        )
        assert r.status_code == 404


def test_startup_recovery():
    uid = f"uid-ext-recovery-{uuid.uuid4().hex}"
    with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
        eval_id, req_doc, _ = _setup_eval_with_all_docs(uid)

    job_ids: list[uuid.UUID] = []

    async def _seed_interrupted_state():
        async with async_session() as s:
            for st in ("queued", "extracting", "analysing"):
                jid = uuid.uuid4()
                job_ids.append(jid)
                s.add(
                    AnalysisJob(
                        id=jid,
                        evaluation_id=uuid.UUID(eval_id),
                        kind="extract_requirements",
                        status=st,
                        step_label="In progress",
                    )
                )
            doc = await s.get(Document, uuid.UUID(req_doc["id"]))
            doc.status = "extracting"
            await s.commit()

        await recover_interrupted_jobs(async_session)

        async with async_session() as s:
            for jid in job_ids:
                j = await s.get(AnalysisJob, jid)
                assert j.status == "failed"
                assert j.error_message == INTERRUPTED_MESSAGE
                assert j.finished_at is not None
            doc = await s.get(Document, uuid.UUID(req_doc["id"]))
            assert doc.status == "error"
            assert doc.error_message == INTERRUPTED_MESSAGE

    asyncio.run(_seed_interrupted_state())


def test_worker_entry_runs_with_only_job_id_and_factories():
    uid = f"uid-ext-direct-worker-{uuid.uuid4().hex}"
    with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
        eval_id, _, _ = _setup_eval_with_all_docs(uid)

    fake_gemini = FakeGeminiClient()

    async def _run_direct():
        async with async_session() as s:
            job = AnalysisJob(
                id=uuid.uuid4(),
                evaluation_id=uuid.UUID(eval_id),
                kind="extract_requirements",
                status="queued",
                step_label="Queued",
            )
            s.add(job)
            await s.commit()
            jid = job.id

        await run_job_by_id(jid, async_session, lambda: fake_gemini)

        async with async_session() as s:
            finished_job = await s.get(AnalysisJob, jid)
            assert finished_job.status == "done"
            assert finished_job.finished_at is not None
            ev = await s.get(Evaluation, uuid.UUID(eval_id))
            assert ev.status == "requirements_review"

    asyncio.run(_run_direct())


def test_worker_module_has_no_fastapi_import():
    worker_path = Path(__file__).resolve().parents[2] / "app" / "services" / "worker.py"
    res = subprocess.run(
        ["grep", "-i", "fastapi", str(worker_path)],
        capture_output=True,
        text=True,
    )
    assert res.returncode == 1, f"worker.py must not reference FastAPI: {res.stdout}"


def test_docx_conversion_dataset_shubh_proposal():
    dataset_dir = _find_dataset_dir()
    shubh_path = dataset_dir / "proposals" / "03_Shubh_Proposal.docx"
    assert shubh_path.exists(), f"Missing dataset file at {shubh_path}"

    docx_bytes = shubh_path.read_bytes()
    pages = extract_docx_pages(docx_bytes)
    assert len(pages) == 1
    assert pages[0].page_number == 1
    text = pages[0].text
    assert "## 1. Service schedule" in text
    assert "## 6. Commercials" in text
    assert " | " in text


def test_reupload_removes_old_pages():
    uid = f"uid-ext-reupload-{uuid.uuid4().hex}"
    fake_gemini = FakeGeminiClient(
        default_pages=[
            PageText(page_number=1, text="Page 1 content"),
            PageText(page_number=2, text="Page 2 content"),
        ]
    )
    app.dependency_overrides[get_job_launcher] = lambda: InlineJobLauncher(
        async_session, lambda: fake_gemini
    )
    try:
        with patch("firebase_admin.auth.verify_id_token", return_value=_fake_token(uid)):
            auth = {"Authorization": "Bearer token"}
            eval_id, req_doc, _ = _setup_eval_with_all_docs(uid, pdf_pages=2)

            resp = client.post(f"/api/evaluations/{eval_id}/extract", headers=auth)
            assert resp.status_code == 202

            old_doc_id = uuid.UUID(req_doc["id"])

            async def _count_pages(doc_id: uuid.UUID) -> int:
                async with async_session() as s:
                    stmt = select(DocumentPage).where(DocumentPage.document_id == doc_id)
                    rows = (await s.execute(stmt)).scalars().all()
                    return len(rows)

            assert asyncio.run(_count_pages(old_doc_id)) == 2

            # Re-upload requirements document -> replaces old document and cascades delete on its pages
            new_pdf = make_valid_pdf_bytes(page_count=1)
            r_reupload = client.post(
                f"/api/evaluations/{eval_id}/documents",
                headers=auth,
                files={"file": ("rfp_v2.pdf", new_pdf, "application/pdf")},
                data={"kind": "requirements"},
            )
            assert r_reupload.status_code == 200
            assert asyncio.run(_count_pages(old_doc_id)) == 0
    finally:
        app.dependency_overrides.pop(get_job_launcher, None)


@pytest.mark.skipif(
    os.getenv("RUN_GEMINI_LIVE") != "1",
    reason="Live Gemini smoke test skipped unless RUN_GEMINI_LIVE=1",
)
def test_live_gemini_smoke_dataset_requirements_pdf():
    dataset_dir = _find_dataset_dir()
    pdf_path = dataset_dir / "requirements" / "01_Requirements_Andheri_Housekeeping.pdf"
    assert pdf_path.exists(), f"Missing dataset PDF at {pdf_path}"

    pdf_bytes = pdf_path.read_bytes()
    gemini_client = create_gemini_client()

    async def _run_live():
        start = time.monotonic()
        pages = await gemini_client.pdf_to_pages(pdf_bytes)
        reqs = await gemini_client.extract_requirements(pages)
        elapsed = time.monotonic() - start
        return pages, reqs, elapsed

    pages, reqs, elapsed = asyncio.run(_run_live())
    total_pages = len(pages)
    total_reqs = len(reqs)
    mandatory_count = sum(1 for r in reqs if r.mandatory)

    print(
        f"\n[LIVE_SMOKE_RESULT] pages={total_pages} requirements={total_reqs} mandatory={mandatory_count} duration_s={elapsed:.2f}"
    )

    assert total_pages >= 1
    assert total_reqs == 8, f"Expected 8 requirements, got {total_reqs}"
    assert mandatory_count == 6, f"Expected 6 mandatory requirements, got {mandatory_count}"
    for r in reqs:
        assert r.source_page is not None, f"Missing source_page on {r.code}"
        assert 1 <= r.source_page <= total_pages, (
            f"source_page {r.source_page} out of range 1..{total_pages}"
        )
