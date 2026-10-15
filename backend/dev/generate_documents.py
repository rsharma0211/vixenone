"""Dev-only document generator for VixenOne demo evaluation packs.

Generates realistic fictional RFQ and proposal documents (PDF and DOCX)
for the Powai, Pantry (BKC), and Thane demo evaluations using ONLY
the factual statements specified in the Phase 6 B.2 prompt.

Excluded from production Docker image via backend/.dockerignore.
"""

from pathlib import Path
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FOOTER_TEXT = "Fictional demo document"
ISSUER_NAME = "Northbridge Properties Pvt Ltd"


def _add_pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#64748b"))
    canvas.drawCentredString(doc.pagesize[0] / 2.0, 24, FOOTER_TEXT)
    canvas.restoreState()


def _create_pdf(target_path: Path, title: str, subtitle: str, sections: list) -> None:
    target_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(target_path),
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#142b45"),
        spaceAfter=4,
    )
    sub_style = ParagraphStyle(
        "DocSub",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#334155"),
        spaceAfter=12,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=6,
    )
    h2_style = ParagraphStyle(
        "DocH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#142b45"),
        spaceBefore=8,
        spaceAfter=4,
    )
    table_cell = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1e293b"),
    )
    table_header = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
    )

    story = [
        Paragraph(title, title_style),
        Paragraph(subtitle, sub_style),
        Spacer(1, 4),
    ]

    for item in sections:
        kind = item.get("kind", "para")
        if kind == "h2":
            story.append(Paragraph(item["text"], h2_style))
        elif kind == "para":
            story.append(Paragraph(item["text"], body_style))
        elif kind == "table":
            raw_data = item["data"]
            col_widths = item.get("widths", None)
            table_data = []
            for row_idx, row in enumerate(raw_data):
                styled_row = []
                for cell in row:
                    st = table_header if row_idx == 0 else table_cell
                    styled_row.append(Paragraph(str(cell), st))
                table_data.append(styled_row)
            t = Table(table_data, colWidths=col_widths)
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#142b45")),
                        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 6))

    doc.build(story, onFirstPage=_add_pdf_footer, onLaterPages=_add_pdf_footer)


def _create_docx(target_path: Path, title: str, subtitle: str, sections: list) -> None:
    target_path.parent.mkdir(parents=True, exist_ok=True)
    doc = DocxDocument()

    # Footer
    footer = doc.sections[0].footer
    footer_p = footer.paragraphs[0]
    footer_p.text = FOOTER_TEXT
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if footer_p.runs:
        footer_p.runs[0].font.size = Pt(8.5)
        footer_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    # Title
    t_p = doc.add_paragraph()
    r = t_p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(20, 43, 69)

    # Subtitle
    s_p = doc.add_paragraph()
    r2 = s_p.add_run(subtitle)
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(51, 65, 85)

    for item in sections:
        kind = item.get("kind", "para")
        if kind == "h2":
            h = doc.add_heading(item["text"], level=2)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
        elif kind == "para":
            p = doc.add_paragraph(item["text"])
            p.paragraph_format.space_after = Pt(4)
        elif kind == "table":
            raw_data = item["data"]
            t = doc.add_table(rows=len(raw_data), cols=len(raw_data[0]))
            t.style = "Table Grid"
            for row_idx, row in enumerate(raw_data):
                for col_idx, cell in enumerate(row):
                    cell_obj = t.cell(row_idx, col_idx)
                    cell_obj.text = str(cell)
                    for prg in cell_obj.paragraphs:
                        for run in prg.runs:
                            run.font.size = Pt(8.5)
                            if row_idx == 0:
                                run.bold = True

    doc.save(str(target_path))


# ==============================================================================
# POWAI — Security Services
# ==============================================================================
def generate_powai_documents(dataset_dir: Path) -> None:
    req_path = dataset_dir / "requirements" / "05_Requirements_Powai_Security.pdf"
    v1_path = dataset_dir / "proposals" / "06_SecureShield_Proposal.pdf"
    v2_path = dataset_dir / "proposals" / "07_Apex_Proposal.docx"

    # Powai Requirements PDF
    _create_pdf(
        req_path,
        title="Request for Proposal — Security Services",
        subtitle=f"Powai Campus · {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Scope of Requirements",
            },
            {
                "kind": "para",
                "text": (
                    f"{ISSUER_NAME} invites proposals from qualified private security agencies "
                    "for round-the-clock physical security coverage across its Powai Campus facility. "
                    "The requirements listed below define the mandatory and operational expectations."
                ),
            },
            {
                "kind": "table",
                "widths": [55, 230, 65, 150],
                "data": [
                    ["Code", "Requirement", "Mandatory", "Acceptance Criteria"],
                    [
                        "REQ-01",
                        "Trained guards 24×7 at all three campus gates (mandatory).",
                        "Yes",
                        "Duty roster confirms 24×7 guard deployment at Gates 1, 2, and 3.",
                    ],
                    [
                        "REQ-02",
                        "Valid private security agency licence for Maharashtra (mandatory).",
                        "Yes",
                        "Active PSARA licence copy for Maharashtra submitted.",
                    ],
                    [
                        "REQ-03",
                        "One security supervisor on site every shift (mandatory).",
                        "Yes",
                        "Designated supervisor deployed on site for every active shift.",
                    ],
                    [
                        "REQ-04",
                        "Visitor log at reception for every entry and exit (mandatory).",
                        "Yes",
                        "Visitor register or digital log maintained at reception for all visits.",
                    ],
                    [
                        "REQ-05",
                        "Fire-safety training within 30 days of deployment.",
                        "No",
                        "Guards complete certified fire-safety training within 30 days.",
                    ],
                    [
                        "REQ-06",
                        "Monthly incident report to the facilities manager.",
                        "No",
                        "Comprehensive incident report submitted monthly.",
                    ],
                ],
            },
        ],
    )

    # SecureShield Proposal PDF
    _create_pdf(
        v1_path,
        title="SecureShield Guards — Technical & Commercial Proposal",
        subtitle=f"Security Services for Powai Campus · Submitted to {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Operational Deployment & Scope",
            },
            {
                "kind": "para",
                "text": (
                    "SecureShield Guards commits to deploying trained guards 24×7 at all three campus gates. "
                    "Continuous vigilance and physical security will be maintained across all campus entry and exit points."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Our agency holds a valid private security agency licence for Maharashtra under registration "
                    "number PSARA-MH-2023-F9876, valid to 2028."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "One dedicated security supervisor is stationed on site every shift to oversee guard operations "
                    "and handle immediate escalations."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "A formal visitor register at reception is strictly maintained for every entry and exit. "
                    "All deployed personnel undergo certified fire-safety training within 30 days of initial deployment."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "A detailed monthly incident report will be submitted directly to the facilities manager at the close of every month."
                ),
            },
            {
                "kind": "h2",
                "text": "2. Commercial Terms",
            },
            {
                "kind": "table",
                "widths": [320, 180],
                "data": [
                    ["Commercial Item", "Monthly Fee (INR)"],
                    ["Monthly base security fee (all 3 gates 24×7, supervisor every shift)", "₹3,20,000"],
                    ["GST @ 18%", "₹57,600"],
                    ["Total Monthly Billing", "₹3,77,600"],
                ],
            },
            {
                "kind": "para",
                "text": (
                    "The monthly rate of ₹3,20,000 + GST includes all guard deployments, on-site supervision every shift, "
                    "visitor register management, fire-safety training within 30 days, and monthly incident reporting."
                ),
            },
        ],
    )

    # Apex Patrol Services DOCX
    _create_docx(
        v2_path,
        title="Apex Patrol Services — Proposal for Security Services",
        subtitle=f"Powai Campus · Prepared for {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Guard Coverage and Operations",
            },
            {
                "kind": "para",
                "text": (
                    "Apex Patrol Services proposes guard coverage with Gates 1–2 manned 24×7, while Gate 3 will have day shift only coverage."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Regarding statutory compliance, our private security agency licence renewal applied for with the controlling authority."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "On-site supervisory coverage includes day supervision; supervisor night shift on request."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "A visitor log maintained at reception for tracking building entries and exits."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Specialised guard fire-safety training is available on request, at additional cost."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Operations will provide a quarterly incident report to client management."
                ),
            },
            {
                "kind": "h2",
                "text": "2. Commercial Quotation",
            },
            {
                "kind": "table",
                "data": [
                    ["Item Description", "Price (INR)"],
                    ["Base security services fee", "₹2,65,000/month + GST"],
                    ["Overtime", "at actuals"],
                    ["Fire-safety training", "on request, at additional cost"],
                ],
            },
            {
                "kind": "para",
                "text": (
                    "The monthly base contract fee is ₹2,65,000/month + GST. "
                    "Any guard overtime beyond scheduled shift hours will be billed at actuals. "
                    "Night shift supervision and fire-safety training are provided on request at additional cost."
                ),
            },
        ],
    )


# ==============================================================================
# PANTRY — Pantry & Cafeteria, BKC
# ==============================================================================
def generate_pantry_documents(dataset_dir: Path) -> None:
    req_path = dataset_dir / "requirements" / "08_Requirements_Pantry_BKC.pdf"
    v1_path = dataset_dir / "proposals" / "09_FreshBite_Proposal.docx"
    v2_path = dataset_dir / "proposals" / "10_Gourmet_Proposal.pdf"

    # Pantry Requirements PDF
    _create_pdf(
        req_path,
        title="Request for Proposal — Pantry & Cafeteria Services",
        subtitle=f"BKC Office · {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Requirements Specifications",
            },
            {
                "kind": "para",
                "text": (
                    f"{ISSUER_NAME} requires a qualified catering and cafeteria partner to operate the corporate pantry "
                    "and cafeteria at its BKC Office location. All catering operations must fulfill the requirements below."
                ),
            },
            {
                "kind": "table",
                "widths": [55, 230, 65, 150],
                "data": [
                    ["Code", "Requirement", "Mandatory", "Acceptance Criteria"],
                    [
                        "REQ-01",
                        "Valid FSSAI licence (mandatory).",
                        "Yes",
                        "Active FSSAI food business operating licence provided.",
                    ],
                    [
                        "REQ-02",
                        "Breakfast and lunch Monday to Friday for up to 150 staff (mandatory).",
                        "Yes",
                        "Daily breakfast and lunch service Mon–Fri covering up to 150 personnel.",
                    ],
                    [
                        "REQ-03",
                        "Daily kitchen hygiene checklist signed by the supervisor (mandatory).",
                        "Yes",
                        "Checklist executed daily and signed by the catering supervisor.",
                    ],
                    [
                        "REQ-04",
                        "Weekly menu shared every Friday.",
                        "No",
                        "Planned menu circulated to facilities management every Friday.",
                    ],
                    [
                        "REQ-05",
                        "Tea and coffee stations refilled twice daily.",
                        "No",
                        "Beverage stations replenished at least twice per day.",
                    ],
                ],
            },
        ],
    )

    # FreshBite Caterers DOCX
    _create_docx(
        v1_path,
        title="FreshBite Caterers — Pantry & Cafeteria Proposal",
        subtitle=f"BKC Office Catering · Submitted to {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Operational Capabilities and Scope",
            },
            {
                "kind": "para",
                "text": (
                    "FreshBite Caterers holds an active food safety licence with FSSAI licence number stated: FSSAI-MH-2023-009876543210."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "We provide full breakfast + lunch Mon–Fri up to 150 staff at the corporate cafeteria. "
                    "All meals are prepared fresh adhering to high quality and nutritional standards."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Our kitchen team maintains a strict daily hygiene checklist signed by supervisor to ensure sanitation."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "A rotational weekly menu every Friday is circulated in advance to the facilities team."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Corporate tea/coffee refilled twice daily across all designated pantry beverage points."
                ),
            },
            {
                "kind": "h2",
                "text": "2. Commercial Quotation",
            },
            {
                "kind": "table",
                "data": [
                    ["Service Component", "Monthly Pricing (INR)"],
                    ["Base cafeteria service fee (breakfast & lunch up to 150 staff)", "₹2,10,000"],
                    ["GST @ 18%", "₹37,800"],
                    ["Total Monthly Fee", "₹2,47,800"],
                ],
            },
            {
                "kind": "para",
                "text": (
                    "Commercial fee is ₹2,10,000/month + GST. "
                    "All culinary raw materials included in the monthly fee. No additional charges apply."
                ),
            },
        ],
    )

    # Gourmet Pantry Co PDF
    _create_pdf(
        v2_path,
        title="Gourmet Pantry Co — Corporate Catering Proposal",
        subtitle=f"Cafeteria & Pantry Services for BKC Office · {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Service Specifications",
            },
            {
                "kind": "para",
                "text": (
                    "Gourmet Pantry Co notes that statutory FSSAI registration in process under reference FSSAI-ACK-2026-00412."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "We offer breakfast + lunch Mon–Sat up to 120 staff. "
                    "Service covers up to 120 staff members on site."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Kitchen cleanliness is audited via our weekly hygiene checklist."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "The cafeteria meal menu shared monthly with the client representative."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Standard tea/coffee refilled once daily, extra refills chargeable."
                ),
            },
            {
                "kind": "h2",
                "text": "2. Financial Proposal",
            },
            {
                "kind": "table",
                "widths": [320, 180],
                "data": [
                    ["Billing Element", "Pricing (INR)"],
                    ["Monthly cafeteria base fee (up to 120 staff)", "₹1,85,000"],
                    ["GST @ 18%", "₹33,300"],
                    ["Total Base Billing", "₹2,18,300"],
                    ["Plates beyond 120", "at actuals"],
                    ["Extra tea/coffee refills", "chargeable"],
                ],
            },
            {
                "kind": "para",
                "text": (
                    "Base pricing is ₹1,85,000/month + GST. "
                    "Plates beyond 120 at actuals will be invoiced monthly. "
                    "Extra tea and coffee refills beyond the once daily refill are chargeable."
                ),
            },
        ],
    )


# ==============================================================================
# THANE — Pest Control AMC
# ==============================================================================
def generate_thane_documents(dataset_dir: Path) -> None:
    req_path = dataset_dir / "requirements" / "11_Requirements_Thane_PestControl.pdf"
    v1_path = dataset_dir / "proposals" / "12_TermiGuard_Proposal.pdf"
    v2_path = dataset_dir / "proposals" / "13_SafeClean_Proposal.docx"

    # Thane Requirements PDF
    _create_pdf(
        req_path,
        title="Request for Proposal — Pest Control AMC",
        subtitle=f"Thane Warehouse · {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Scope of Pest Mitigation Requirements",
            },
            {
                "kind": "para",
                "text": (
                    f"{ISSUER_NAME} invites proposals for an Annual Maintenance Contract (AMC) "
                    "for complete pest control at its Thane Warehouse premises."
                ),
            },
            {
                "kind": "table",
                "widths": [55, 230, 65, 150],
                "data": [
                    ["Code", "Requirement", "Mandatory", "Acceptance Criteria"],
                    [
                        "REQ-01",
                        "Monthly general pest control of the whole warehouse (mandatory).",
                        "Yes",
                        "General pest management carried out monthly across the entire warehouse.",
                    ],
                    [
                        "REQ-02",
                        "Rodent bait stations on the perimeter, checked fortnightly (mandatory).",
                        "Yes",
                        "Perimeter bait stations inspected and serviced on a fortnightly cadence.",
                    ],
                    [
                        "REQ-03",
                        "Annual termite treatment (mandatory).",
                        "Yes",
                        "Complete termite preventive/curative treatment conducted annually.",
                    ],
                    [
                        "REQ-04",
                        "Only government-approved chemicals; safety data sheets shared (mandatory).",
                        "Yes",
                        "Certified chemicals used exclusively; manufacturer safety data sheets provided.",
                    ],
                    [
                        "REQ-05",
                        "Service report after each visit.",
                        "No",
                        "Service activity report submitted immediately after every service visit.",
                    ],
                    [
                        "REQ-06",
                        "Emergency call-out within 24 hours.",
                        "No",
                        "Response and site visit provided within 24 hours of an emergency notice.",
                    ],
                ],
            },
        ],
    )

    # TermiGuard AMC PDF
    _create_pdf(
        v1_path,
        title="TermiGuard AMC — Warehouse Pest Management Proposal",
        subtitle=f"Thane Warehouse · Submitted to {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Pest Management Scope and Methodology",
            },
            {
                "kind": "para",
                "text": (
                    "TermiGuard AMC guarantees monthly general treatment covering the full warehouse storage and dispatch zones."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Tamper-proof rodent stations checked fortnightly along the entire facility perimeter."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Comprehensive annual termite treatment included in the contract scope at no added cost."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Our technicians utilize only approved chemicals with safety data sheets submitted before treatment."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "A formal technician report after each visit will be handed to the warehouse supervisor."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "In the event of acute pest sighting, guaranteed emergency within 24 hours."
                ),
            },
            {
                "kind": "h2",
                "text": "2. Financial Proposal",
            },
            {
                "kind": "table",
                "widths": [320, 180],
                "data": [
                    ["AMC Component", "Amount (INR)"],
                    ["Monthly general pest control fee", "₹15,000/month + GST"],
                    ["Billing cycle", "billed quarterly (₹45,000 + GST)"],
                    ["Annual termite treatment", "Included"],
                    ["Safety data sheets", "Included"],
                ],
            },
            {
                "kind": "para",
                "text": (
                    "Pricing is ₹15,000/month + GST, billed quarterly. "
                    "Annual termite treatment, fortnightly rodent servicing, safety data sheets, and 24-hour emergency response are fully included."
                ),
            },
        ],
    )

    # SafeClean Disinfection DOCX
    _create_docx(
        v2_path,
        title="SafeClean Disinfection — Commercial Pest AMC Proposal",
        subtitle=f"Thane Warehouse Facility · Prepared for {ISSUER_NAME}",
        sections=[
            {
                "kind": "h2",
                "text": "1. Operational Scope",
            },
            {
                "kind": "para",
                "text": (
                    "SafeClean Disinfection provides scheduled monthly general treatment across the warehouse premises."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Perimeter rodent stations checked monthly by our field team."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Subterranean termite treatment excluded, quoted separately upon specialized survey."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "We use approved chemicals, safety sheets on request."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "Field staff will supply a service report after each visit."
                ),
            },
            {
                "kind": "para",
                "text": (
                    "For unexpected pest escalations, emergency call-outs are handled on a best effort basis."
                ),
            },
            {
                "kind": "h2",
                "text": "2. Commercial Schedule",
            },
            {
                "kind": "table",
                "data": [
                    ["Service Element", "Rate (INR)"],
                    ["Monthly general pest management fee", "₹12,500/month + GST"],
                    ["Termite treatment", "Excluded (quoted separately)"],
                    ["Emergency call-outs", "best effort"],
                ],
            },
            {
                "kind": "para",
                "text": (
                    "The contract rate is ₹12,500/month + GST. "
                    "Includes general monthly treatment and monthly rodent checks. "
                    "Termite treatment is excluded and quoted separately. Safety sheets on request."
                ),
            },
        ],
    )


def generate_all_documents(dataset_dir: Path) -> None:
    print(f"Generating documents into {dataset_dir}...")
    generate_powai_documents(dataset_dir)
    generate_pantry_documents(dataset_dir)
    generate_thane_documents(dataset_dir)
    print("All documents generated successfully.")


if __name__ == "__main__":
    import sys

    base = Path(__file__).resolve().parent.parent / "dataset"
    if len(sys.argv) > 1:
        base = Path(sys.argv[1])
    generate_all_documents(base)
